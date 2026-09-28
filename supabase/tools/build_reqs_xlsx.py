import re, collections, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import reqs_data as rd
R = rd.R
BR = [
("BR-001","Enterprise","One controlled system of record for hotel operations.","Must","R1"),
("BR-002","Enterprise","Property-level configuration and future multi-property operation.","Must","R1"),
("BR-003","Enterprise","A single guest identity across reservations and stays.","Must","R1"),
("BR-004","Reservations","Sellable room inventory for any requested stay period.","Must","R1"),
("BR-005","Reservations","Prevent conflicting physical room assignments.","Must","R1"),
("BR-006","Front Office","Arrivals, departures, in-house guests and room assignments from a focused interface.","Must","R1"),
("BR-007","Front Office","Walk-in, direct, corporate, group and future channel bookings. Reworded (A3): corporate is a booking source, a free-text company name and a corporate rate plan. No company entity.","Must","R1"),
("BR-008","Guest","Guest history accessible to authorised staff across stays.","Must","R1"),
("BR-009","Rooms","Commercial availability kept apart from operational room condition.","Must","R1"),
("BR-010","Housekeeping","Check-out triggers the room turnover workflow.","Must","R1"),
("BR-011","Housekeeping","Only rooms meeting readiness rules are normally assignable for arrival.","Must","R1"),
("BR-012","Maintenance","Maintenance issues that affect sellability can block the room.","Must","R1/R2"),
("BR-013","Finance","All financial activity is a traceable folio transaction.","Must","R1"),
("BR-014","Finance","Financial transactions are never silently deleted.","Must","R1"),
("BR-015","Finance","Refunds, reversals, discounts and adjustments keep an audit trail.","Must","R1"),
("BR-016","Finance","Reconcile payments by cashier, shift, method and business date.","Must","R2 (method and date in R1)"),
("BR-017","Security","Users act only within their role and property scope.","Must","R1"),
("BR-018","Security","Sensitive actions record actor, time, action, before and after values, and reason.","Must","R1"),
("BR-019","Reporting","Common KPI definitions for occupancy, ADR, RevPAR, revenue and status.","Must","R1"),
("BR-020","Reporting","Reports trace to the underlying records.","Must","R1"),
("BR-021","F&B","Restaurant and bar orders posting room charges.","Should","R2"),
("BR-022","Inventory","Stock receipts, issues, transfers, wastage and counts.","Should","R2"),
("BR-023","Procurement","Request, approval, PO and receipt controls.","Should","R2"),
("BR-024","Commercial","Corporate and group business as structured commercial accounts.","Should","R2 (R1 covers source, name, rate plan, group record)"),
("BR-025","Events","Venue, services, accommodation and billing for events.","Should","R2"),
("BR-026","Workforce","Staff records and schedules for assignment.","Should","R2"),
("BR-027","Digital","Guest self-service for booking and stay.","Could","R3"),
("BR-028","Distribution","Channel sync of inventory and rates.","Could","R3"),
("BR-029","Loyalty","Loyalty and guest segmentation.","Could","R3"),
("BR-030","AI","AI-assisted analysis within authorised controls.","Could","R4"),
]
DEC = [
("A1","Multi-tenant product for many hotels.","Confirmed"),
("A2","Business date rolls automatically in R1. Night audit and cashier shift come in R2.","Confirmed"),
("A3","No Company entity. Corporate is a source, a free-text company name and a corporate rate plan. BR-007 reworded.","Confirmed"),
("A10","Nigeria first. NGN base currency. Multi-currency with recorded FX rate. Payments are taken outside the system.","Confirmed"),
("A11","Tax seed (10% service charge, 7.5% VAT on net plus prior, state levy 0%) is unverified.","Open: accountant to confirm"),
("A12","Online only in R1. Client ids and idempotency keys keep the door open for offline.","Confirmed"),
("A13","Self-service signup with a 30-day trial. Billing integration is R1b. Pilots are invoiced by hand.","Confirmed"),
("A14","Cloud hosting. Provider and region not chosen. Nearest Supabase region is Frankfurt.","Open: pick region"),
("A16","Plans are configuration. Proposed: 30-day trial and room-count tiers (starter 20, standard 60, pro 150, enterprise unlimited). Unpriced.","Open: pick plans and prices"),
("DAY-006","Same-day arrival and departure: charge a night, a day-use rate, or nothing.","Open: decide"),
("Fees","Cancellation and no-show fee charge codes are non-taxable in the seed.","Open: accountant to confirm"),
("Ops","Enable pg_cron in Supabase and schedule app.run_due_rollovers() every 5 minutes. Run the subscription clock daily.","Open: setup step"),
]
GAPS = [
("PLT-010","Break-glass support access","Not built","Schema and RPC for time-boxed access with reason, tenant notice and audit."),
("ADM-009","Data exports","Not built","Export RPCs or Edge Function with permission check and audit entry."),
("ADM-001","Invite staff","Edge Function","Function creates the auth user, then calls register_invited_user."),
("RES-012","Group booking UI","UI","Backend done. UI to create several one-room reservations under one group."),
("GST-006","Feedback and complaints UI","UI","Backend done. UI needed."),
("RPT-004","Report views","Views","Add views for cancellation rate, ALOS and outstanding balances."),
("PLT-012","Subscription billing","R1b","Paystack or Flutterwave integration."),
("FIN-016","Cashier shifts","R1b/R2","Shift open, close and reconcile."),
("DAY-005","Night audit gate","R1b/R2","Audit must pass before the date rolls."),
]
hdr = ["ID","Module","Screens","Requirement","BR trace","Priority","Release","Permissions","Transition / rule refs","Given","When","Then","Built by","Verified by (tests)","Status"]
navy=PatternFill("solid",fgColor="1F3A5F"); grey=PatternFill("solid",fgColor="F2F2F2")
thin=Side(style="thin",color="BFBFBF"); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
stat_fill={"Built+tested":"D9EAD3","Built (UI or Edge Function work remains)":"FFF2CC","R1b":"D9E2F3","Not built":"F4CCCC","Decision needed":"FCE5CD"}
def sheet(ws, head, rows, widths, statcol=None):
    ws.append(head)
    for c in ws[1]:
        c.font=Font(bold=True,color="FFFFFF"); c.fill=navy; c.alignment=Alignment(wrap_text=True,vertical="center"); c.border=bd
    for r in rows: ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment=Alignment(wrap_text=True,vertical="top"); c.border=bd
        if statcol is not None:
            v=row[statcol].value
            if v in stat_fill: row[statcol].fill=PatternFill("solid",fgColor=stat_fill[v])
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes="B2"; ws.auto_filter.ref=ws.dimensions
wb=Workbook()
ws=wb.active; ws.title="README"
cnt=collections.Counter(x[14] for x in R)
lines=[("HMS R1 Requirements",),("Version 0.5, 28 Sep 2026",),(),
("Sheets",),("R1 Requirements","One row per requirement, with Given, When, Then."),("BR Coverage","Each business requirement mapped to requirement ids."),
("Decisions","Confirmed decisions and open questions."),("Build Gaps","What is not yet built."),(),
("Counts",),("Total requirements",len(R))]+[(k,v) for k,v in cnt.items()]+[(),
("How to read a row",),("Given / When / Then","The acceptance criterion. A tester can run it as written."),
("Permissions","Permission keys checked by the database. Grant level Y, L (limit) or O (own)."),
("Verified by","Test file number in supabase/tests. Run with supabase/tests/run.sh."),
("Built by","Database function or object that implements the requirement.")]
for l in lines: ws.append(list(l))
ws["A1"].font=Font(bold=True,size=16)
for r in (4,10,(15+len(cnt))): ws.cell(r,1).font=Font(bold=True)
ws.column_dimensions["A"].width=26; ws.column_dimensions["B"].width=80
sheet(wb.create_sheet("R1 Requirements"),hdr,R,[10,16,18,45,14,9,9,26,28,38,38,48,26,10,20],statcol=14)
idx=collections.defaultdict(list)
for x in R:
    for b in re.findall(r'BR-\d+',x[4]): idx[b].append(x[0])
rows=[]
for b,dom,txt,pr,rel in BR:
    ids=idx.get(b,[])
    cov="Covered in R1" if ids and rel.startswith("R1") else ("Partly covered in R1" if ids else "Later release")
    rows.append((b,dom,txt,pr,rel,", ".join(ids),len(ids),cov))
sheet(wb.create_sheet("BR Coverage"),["BR","Domain","Business requirement","Priority","Release","Requirement ids","Count","Coverage"],rows,[9,14,60,9,26,50,8,20])
sheet(wb.create_sheet("Decisions"),["Ref","Decision or question","State"],DEC,[10,90,28])
sheet(wb.create_sheet("Build Gaps"),["ID","Item","Type","Work remaining"],GAPS,[10,30,14,80])
out="/home/claude/hms/docs/HMS_R1_Requirements.xlsx"
wb.save(out); print(out, len(R))
